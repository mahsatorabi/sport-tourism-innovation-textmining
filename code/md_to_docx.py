"""Convert Article1_Methods_Results.md to a Q1-style DOCX."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
MD = ROOT / "Article1_Methods_Results.md"
OUT = ROOT / "Article1_Methods_Results.docx"

INLINE_RE = re.compile(
    r"(\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+?`)"
)


def set_run_font(run, bold=None, italic=None, size=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if size is not None:
        run.font.size = size


def add_formatted_para(doc, line, list_style=None):
    p = doc.add_paragraph(style=list_style) if list_style else doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.5
    for part in INLINE_RE.split(line):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_run_font(run, bold=True)
        elif part.startswith("*") and part.endswith("*"):
            run = p.add_run(part[1:-1])
            set_run_font(run, italic=True)
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            set_run_font(run, italic=True)
        else:
            run = p.add_run(part)
            set_run_font(run)
    return p


def clean_cell(val: str) -> str:
    val = re.sub(r"\*\*([^*]+)\*\*", r"\1", val)
    val = re.sub(r"\*([^*]+)\*", r"\1", val)
    val = val.replace("`", "")
    return val.strip()


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(clean_cell(h))
        set_run_font(run, bold=True, size=Pt(10))
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(clean_cell(val))
            set_run_font(run, size=Pt(10))
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
    doc.add_paragraph()


def main():
    text = MD.read_text(encoding="utf-8")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.5

    for level, size in ((1, 16), (2, 14), (3, 12)):
        hs = doc.styles[f"Heading {level}"]
        hs.font.name = "Times New Roman"
        hs.font.color.rgb = RGBColor(0, 0, 0)
        hs.font.size = Pt(size)
        hs._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line or line.strip() == "---":
            i += 1
            continue

        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=1)
            for run in p.runs:
                set_run_font(run, bold=True, size=Pt(16))
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_heading(line[3:].strip(), level=2)
            for run in p.runs:
                set_run_font(run, bold=True, size=Pt(14))
            i += 1
            continue
        if line.startswith("### ") or line.startswith("#### "):
            content = line.lstrip("#").strip()
            p = doc.add_heading(content, level=3)
            for run in p.runs:
                set_run_font(run, bold=True, size=Pt(12))
            i += 1
            continue

        if line.startswith("|") and i + 1 < len(lines) and re.match(
            r"^\|[\s\-:|]+\|$", lines[i + 1].strip()
        ):
            headers = [c.strip() for c in line.strip().strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, headers, rows)
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", line)
        if m:
            add_formatted_para(doc, m.group(2), list_style="List Number")
            i += 1
            continue

        if line.startswith("- "):
            add_formatted_para(doc, line[2:], list_style="List Bullet")
            i += 1
            continue

        if line.startswith("**") and line.endswith("**") and len(line) < 140:
            p = doc.add_paragraph()
            run = p.add_run(line[2:-2])
            set_run_font(run, bold=True, italic=True)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        add_formatted_para(doc, line)
        i += 1

    doc.save(OUT)
    print(f"Saved: {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
